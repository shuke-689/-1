# -*- coding: utf-8 -*-
"""达人名单输出模块（collect.py / collect_30.py 共用）。

产出三份：
  darens<tag>.json   全字段明细（供 wechat_add.py 消费）
  darens<tag>.csv    清洗后的可读表（含跳过原因、带货品牌占比等诊断列）
  darens<tag>.xlsx   Excel 登记表 —— 用户指定 7 列：
                     达人名称 / 达人微信 / 达人粉丝 / 销售总额 /
                     直播结算总额 / 视频结算总额 / 图文结算总额
openpyxl 缺失时只出 json + csv，不报错。
"""
import csv
import json
import os

XLSX_HEAD = ["达人名称", "达人微信", "达人粉丝", "销售总额",
             "直播结算总额", "视频结算总额", "图文结算总额"]

CSV_HEAD = ["昵称", "联系方式", "类型", "粉丝数", "地区", "来源类目", "主推类目", "内容类型",
            "达人等级", "结算总额(下限)", "结算总额(上限)", "跳过原因",
            "带货件数", "最大品牌", "最大品牌占比", "带货店铺", "uid"]


def fmt_range(d):
    """把 {'low':x,'high':y} 格式化成 'x-y'；无数据返回空串。"""
    if not isinstance(d, dict):
        return ""
    lo, hi = d.get("low"), d.get("high")
    if lo is None and hi is None:
        return ""
    return "%s-%s" % (lo if lo is not None else "", hi if hi is not None else "")


def xlsx_rows(records):
    """按用户要求的 7 列组织 Excel 行。达人微信只填微信号（手机号不算）。"""
    out = []
    for r in records:
        wx = r.get("contact") if r.get("contact_type") == "微信" else ""
        fans = r.get("fans")
        out.append([r.get("nickname") or "", wx or "",
                    fans if fans is not None else "",
                    fmt_range(r.get("settle_total")),
                    fmt_range(r.get("settle_live")),
                    fmt_range(r.get("settle_video")),
                    fmt_range(r.get("settle_image_text"))])
    return out


def write_xlsx(records, path, log=print):
    """写 Excel。openpyxl 不可用时返回 None。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as e:
        log("  ! 未安装 openpyxl（%s）-> 跳过 xlsx，只出 CSV" % str(e)[:60])
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "达人名单"
    ws.append(XLSX_HEAD)
    for c in range(1, len(XLSX_HEAD) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in xlsx_rows(records):
        ws.append(row)
    for i, w in enumerate([26, 20, 12, 16, 16, 16, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def write_csv(records, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(CSV_HEAD)
        for r in records:
            wr.writerow([r.get("nickname", ""), r.get("contact", ""), r.get("contact_type", ""),
                         r.get("fans", ""), r.get("city", ""), r.get("src_cate", ""),
                         " / ".join(r.get("main_cate") or []),
                         " / ".join(r.get("content_type") or []),
                         r.get("level", ""), r.get("sale_low", ""), r.get("sale_high", ""),
                         r.get("skip_reason", ""), r.get("shop_rows", ""),
                         r.get("top_brand", ""), r.get("top_brand_ratio", ""),
                         " / ".join(r.get("shops") or []),
                         r.get("uid", "")])
    return path


def write_outputs(records, out_dir, tag="", log=print):
    """统一输出 json / csv / xlsx，返回 (json_path, csv_path, xlsx_path)。"""
    os.makedirs(out_dir, exist_ok=True)
    json_path = os.path.join(out_dir, "darens%s.json" % tag)
    csv_path = os.path.join(out_dir, "darens%s.csv" % tag)
    xlsx_path = os.path.join(out_dir, "darens%s.xlsx" % tag)

    json.dump(records, open(json_path, "w", encoding="utf-8"),
              ensure_ascii=False, indent=2)
    write_csv(records, csv_path)
    xp = write_xlsx(records, xlsx_path, log=log)

    log("  JSON: %s（%d 条）" % (json_path, len(records)))
    log("  CSV : %s" % csv_path)
    log("  XLSX: %s" % (xp if xp else "（未生成）"))
    return json_path, csv_path, xp
